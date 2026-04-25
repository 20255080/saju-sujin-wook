"""
data_loader.py - 자료 로더

data/ 폴더에서 각 사람의 사주 자료(PDF, xlsx, docx)를 읽어서
하나의 통합 텍스트로 만들어줍니다.
"""

from pathlib import Path
import pdfplumber
from openpyxl import load_workbook
from docx import Document

DATA_DIR = Path(__file__).parent / "data"

PEOPLE = ["김수진", "신성욱"]


def extract_pdf(filepath):
    """PDF에서 페이지별 텍스트 추출"""
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(f"[페이지 {i}]\n{page_text}")
    return "\n\n".join(text_parts)


def extract_xlsx(filepath):
    """xlsx에서 시트별 표 추출 (각 행을 ' | ' 로 구분)"""
    wb = load_workbook(filepath, data_only=True)
    text_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = " | ".join(cells).strip(" |")
            if row_text:
                rows_text.append(row_text)
        if rows_text:
            text_parts.append(f"[시트: {sheet_name}]\n" + "\n".join(rows_text))
    return "\n\n".join(text_parts)


def extract_docx(filepath):
    """docx에서 단락별 텍스트 추출"""
    doc = Document(filepath)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def load_person_data(person_name):
    """
    특정 사람의 모든 자료를 통합 텍스트로 반환.
    파일명에 person_name이 들어있으면 매칭.
    """
    if not DATA_DIR.exists():
        return ""
    
    parts = []
    for file in sorted(DATA_DIR.iterdir()):
        if not file.is_file():
            continue
        if person_name not in file.name:
            continue
        
        try:
            if file.suffix.lower() == ".pdf":
                content = extract_pdf(file)
            elif file.suffix.lower() in [".xlsx", ".xls"]:
                content = extract_xlsx(file)
            elif file.suffix.lower() == ".docx":
                content = extract_docx(file)
            else:
                continue
            
            if content.strip():
                parts.append(f"\n\n========== [참조 자료: {file.name}] ==========\n\n{content}")
        except Exception as e:
            print(f"  ⚠ {file.name} 읽기 실패: {e}")
    
    return "".join(parts)


def load_all_data():
    """앱 시작 시 모든 사람의 데이터 로드"""
    result = {}
    print(f"\n자료 로딩 시작 ({DATA_DIR})...")
    if not DATA_DIR.exists():
        print(f"  ⚠ data/ 폴더가 없습니다. 빈 자료로 진행합니다.")
        return {p: "" for p in PEOPLE}
    
    for person in PEOPLE:
        content = load_person_data(person)
        result[person] = content
        char_count = len(content)
        if char_count > 0:
            print(f"  ✓ {person}: {char_count:,}자 로드됨")
        else:
            print(f"  · {person}: 자료 없음")
    return result


if __name__ == "__main__":
    # 단독 실행 시 테스트
    data = load_all_data()
    print("\n=== 미리보기 ===")
    for name, content in data.items():
        preview = content[:300] if content else "(비어있음)"
        print(f"\n--- {name} ({len(content):,}자) ---")
        print(preview)
